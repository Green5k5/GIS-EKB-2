#!/usr/bin/env python3
"""
Привязка сканов ведомостей к усадьбам (поле scanUrl в js/data-records.js).

Три способа - по тому, в каком виде пришли сканы:

1) Папка с вырезками, по файлу на усадьбу (как у Екатеринбурга:
   « 1.png», «12.jpg», «215_1.png», «117 (2).png» - номер в начале имени):

     python3 tools/attach_scans.py --settlement Уктус --dir assets/scans/uktus

   Если исходники тяжёлые, их можно сразу пережать в WebP (нужен Pillow):

     python3 tools/attach_scans.py --settlement Уктус \\
         --src ~/Загрузки/уктус_сканы --dir assets/scans/uktus --optimize

2) Таблица «номер усадьбы → ссылка» (например, ссылки Яндекс Диска):
   CSV или .xlsx с колонками num и url. Подходит и присланный список
   «Сканы - чего не хватает.xlsx» после заполнения колонки ссылок:

     python3 tools/attach_scans.py --settlement Уктус --csv uktus_links.csv

3) Только целые листы ведомости (без вырезок по усадьбам). CSV с колонками
   sheet и url, где sheet - номер листа, как в поле «Источник»:
   «112», «112 об.», «112а», «112а об.». Каждая усадьба получит скан листа,
   на котором она записана:

     python3 tools/attach_scans.py --settlement Уктус --sheets uktus_sheets.csv

По умолчанию скрипт ничего не меняет, а только показывает, что сделает.
Чтобы записать изменения, добавьте --write (рядом сохранится data-records.js.bak).
Уже заполненные ссылки не трогаются, если не указать --overwrite.
--report file.csv сохраняет список усадеб, для которых скана так и нет.
"""
import argparse, csv, io, json, os, re, shutil, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG_EXT = (".png", ".jpg", ".jpeg", ".webp", ".gif", ".tif", ".tiff")
NUM_RE = re.compile(r"^\s*(\d+)\s*([a-zа-я])?(?=[\s_.(\-]|$)", re.I)
SHEET_RE = re.compile(r"Л\.\s*(.+?)\.?\s*$")


JS_DIR = os.path.join(ROOT, "js")
# Записи усадеб: js/data-records.js (с 01.10.2026), раньше - js/data.js
RECORDS = os.path.join(JS_DIR, "data-records.js")
if not os.path.exists(RECORDS):
    RECORDS = os.path.join(JS_DIR, "data.js")


def _find_var(text, name):
    m = re.search(r"^var " + name + r" = ", text, re.M)
    if not m:
        return None
    value, end = json.JSONDecoder().raw_decode(text, m.end())
    return value, m.end(), end


def read_var(name, *files):
    """Значение `var <name> = …;` из первого js-файла, где оно есть."""
    for fn in files or ("data-records.js", "data-polygons.js", "data.js"):
        path = os.path.join(JS_DIR, fn)
        if os.path.exists(path):
            found = _find_var(open(path, encoding="utf-8").read(), name)
            if found:
                return found[0]
    return None


def write_records(data):
    """Перезаписывает allData, сохраняя оформление файла (читаемое или в одну строку)."""
    text = open(RECORDS, encoding="utf-8").read()
    _, start, end = _find_var(text, "allData")
    pretty = "\n" in text[start:end]
    body = json.dumps(data, ensure_ascii=False, indent=2) if pretty \
        else json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    shutil.copyfile(RECORDS, RECORDS + ".bak")
    with open(RECORDS, "w", encoding="utf-8", newline="\n") as f:
        f.write(text[:start] + body + text[end:])


def base_num(num):
    """«12» → «12», «12.0» → «12», «747a» → «747a»."""
    s = str(num).strip()
    return s[:-2] if s.endswith(".0") else s


def sheet_of(source):
    m = SHEET_RE.search(source or "")
    return norm_sheet(m.group(1)) if m else ""


def norm_sheet(s):
    s = str(s).lower().replace("ё", "е")
    s = re.sub(r"^\s*л\.?\s*", "", s)
    s = re.sub(r"об\.?", "об", s)
    s = re.sub(r"\s+", " ", s).strip(" .")
    return s


ALIASES = {
    "num": ("num", "№", "№ усадьбы", "номер"),
    "url": ("url", "новая ссылка на скан (заполнить)", "ссылка на скан", "ссылка", "скан"),
    "sheet": ("sheet", "лист"),
    "settlement": ("settlement", "поселение"),
}


def read_csv(path, keys):
    """CSV (; , или таб) или .xlsx; колонки ищутся по названиям из ALIASES."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        tables = [[["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
                  for ws in wb.worksheets]
    else:
        raw = open(path, encoding="utf-8-sig").read()
        first = raw.splitlines()[0] if raw else ""
        tables = [list(csv.reader(io.StringIO(raw), delimiter=max(";,\t", key=first.count)))]
    # лист и строка заголовка - первые, где нашлись все нужные колонки
    found = None
    for table in tables:
        for h, header in enumerate(table[:20]):
            low = [c.strip().lower() for c in header]
            cols = {}
            for k in list(keys) + ["settlement"]:
                for alias in ALIASES.get(k, (k,)):
                    if alias in low:
                        cols[k] = low.index(alias)
                        break
            if all(k in cols for k in keys):
                found = (table, h, cols)
                break
        if found:
            break
    if not found:
        sys.exit(f"В {path} не нашлось колонок: {', '.join(keys)}")
    table, h, cols = found
    rows = []
    for r in table[h + 1:]:
        row = {k: (r[i].strip() if i < len(r) else "") for k, i in cols.items()}
        if row.get("num", "x") or row.get("sheet", "x"):
            rows.append(row)
    return rows


def files_by_num(folder):
    """Номер усадьбы → имя файла. Если на номер несколько файлов, берётся
    вариант с пометкой «_1»/«(2)» (обычно это исправленная версия), а в
    отчёт попадает предупреждение."""
    found, dup, bad = {}, {}, []
    for fn in sorted(os.listdir(folder)):
        if not fn.lower().endswith(IMG_EXT):
            continue
        m = NUM_RE.match(fn)
        if not m:
            bad.append(fn)
            continue
        key = m.group(1) + (m.group(2) or "").lower()
        if key in found:
            dup.setdefault(key, [found[key]]).append(fn)
            if len(fn) > len(found[key]):
                found[key] = fn
        else:
            found[key] = fn
    return found, dup, bad


def optimize(src_dir, dst_dir, max_side, quality):
    from PIL import Image
    os.makedirs(dst_dir, exist_ok=True)
    done = {}
    # «12.png» раньше «12_1.png»: исправленная версия с пометкой пишется последней и остаётся
    for fn in sorted(os.listdir(src_dir), key=lambda s: (len(s), s)):
        if not fn.lower().endswith(IMG_EXT):
            continue
        m = NUM_RE.match(fn)
        if not m:
            print(f"  ! пропущен, в имени нет номера усадьбы: {fn}")
            continue
        out = m.group(1) + (m.group(2) or "").lower() + ".webp"
        if out in done:
            print(f"  ! {done[out]} и {fn} - один номер; остаётся {fn}")
        done[out] = fn
        im = Image.open(os.path.join(src_dir, fn))
        im.load()
        if im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        k = max_side / max(im.size)
        if k < 1:
            im = im.resize((round(im.width * k), round(im.height * k)), Image.LANCZOS)
        im.save(os.path.join(dst_dir, out), "WEBP", quality=quality, method=6)
        print(f"  {fn} → {out}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--settlement", required=True, help="Екатеринбург, Нижне-Исетск или Уктус")
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--dir", help="папка сайта со сканами, по файлу на усадьбу (например assets/scans/uktus)")
    g.add_argument("--csv", "--table", dest="csv", help="CSV или .xlsx с колонками num (или «№») и url (или «Новая ссылка на скан (заполнить)»)")
    g.add_argument("--sheets", help="CSV с колонками sheet и url (скан целого листа)")
    ap.add_argument("--src", help="с --dir и --optimize: откуда брать исходные файлы")
    ap.add_argument("--optimize", action="store_true", help="пережать в WebP (нужен Pillow)")
    ap.add_argument("--max-side", type=int, default=2400)
    ap.add_argument("--quality", type=int, default=80)
    ap.add_argument("--overwrite", action="store_true", help="заменять уже заполненные ссылки")
    ap.add_argument("--write", action="store_true", help="записать изменения в js/data-records.js")
    ap.add_argument("--report", help="сохранить CSV со списком усадеб без скана")
    a = ap.parse_args()

    data = read_var("allData")
    recs = [d for d in data if d.get("settlement") == a.settlement]
    if not recs:
        sys.exit(f"Нет записей с поселением «{a.settlement}». "
                 f"Есть: {', '.join(sorted({d.get('settlement') for d in data}))}")

    link = {}          # id записи → ссылка
    notes = []
    if a.dir:
        folder = os.path.join(ROOT, a.dir) if not os.path.isabs(a.dir) else a.dir
        if a.optimize:
            if not a.src:
                sys.exit("С --optimize укажите --src - папку с исходными файлами")
            print(f"Пережимаю {a.src} → {folder}")
            optimize(a.src, folder, a.max_side, a.quality)
        if not os.path.isdir(folder):
            sys.exit(f"Нет папки {folder}")
        rel = os.path.relpath(folder, ROOT).replace(os.sep, "/")
        found, dup, bad = files_by_num(folder)
        for k, fns in dup.items():
            notes.append(f"на №{k} несколько файлов: {', '.join(fns)} - взят {found[k]}")
        for fn in bad:
            notes.append(f"в имени файла нет номера усадьбы: {fn}")
        used = set()
        for d in recs:
            n = base_num(d["num"]).lower()
            fn = found.get(n) or found.get(re.sub(r"[a-zа-я]$", "", n))
            if fn:
                link[d["id"]] = f"{rel}/{fn}"
                used.add(fn)
        for fn in sorted(set(found.values()) - used):
            notes.append(f"файл не подошёл ни к одной усадьбе: {fn}")
    elif a.csv:
        by_num = {}
        for r in read_csv(a.csv, ["num", "url"]):
            if r.get("settlement") and r["settlement"] != a.settlement:
                continue
            if r["num"] and r["url"]:
                by_num[base_num(r["num"]).lower()] = r["url"]
        for d in recs:
            n = base_num(d["num"]).lower()
            u = by_num.pop(n, None) or by_num.get(re.sub(r"[a-zа-я]$", "", n))
            if u:
                link[d["id"]] = u
        for n in by_num:
            notes.append(f"в CSV есть №{n}, а в данных такой усадьбы нет")
    else:
        by_sheet = {norm_sheet(r["sheet"]): r["url"] for r in read_csv(a.sheets, ["sheet", "url"]) if r["url"]}
        for d in recs:
            u = by_sheet.get(sheet_of(d.get("source")))
            if u:
                link[d["id"]] = u
        known = {sheet_of(d.get("source")) for d in recs}
        for s in by_sheet:
            if s not in known:
                notes.append(f"лист «{s}» из CSV не встречается в источниках записей")

    changed = kept = 0
    for d in recs:
        u = link.get(d["id"])
        if not u:
            continue
        if d.get("scanUrl") and d["scanUrl"] != u and not a.overwrite:
            kept += 1
            continue
        if d.get("scanUrl") != u:
            d["scanUrl"] = u
            changed += 1

    missing = [d for d in recs if not d.get("scanUrl")]
    print(f"\n{a.settlement}: записей {len(recs)}, найдено сканов {len(link)}, "
          f"будет изменено {changed}" + (f", оставлено старых ссылок {kept} (см. --overwrite)" if kept else ""))
    print(f"Без скана останется: {len(missing)}")
    if missing:
        nums = [base_num(d["num"]) for d in missing]
        print("  №: " + ", ".join(nums[:60]) + (" …" if len(nums) > 60 else ""))
    for n in notes:
        print("  ! " + n)

    if a.report:
        with open(a.report, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Поселение", "№", "Владелец", "Источник", "Лист"])
            for d in missing:
                w.writerow([d["settlement"], base_num(d["num"]),
                            " ".join(x for x in (d.get("surname"), d.get("name"), d.get("patronymic")) if x),
                            d.get("source", ""), sheet_of(d.get("source"))])
        print(f"Список без скана: {a.report}")

    if a.write and changed:
        write_records(data)
        rel = os.path.relpath(RECORDS, ROOT).replace(os.sep, "/")
        print(f"{rel} обновлён (копия прежней версии - {rel}.bak). Поменяйте метку версии этого файла в index.html.")
    elif changed:
        print("\nЭто пробный прогон. Чтобы записать, добавьте --write")


if __name__ == "__main__":
    main()
