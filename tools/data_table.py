#!/usr/bin/env python3
"""
Данные усадеб ⇄ таблица. Минимальная «админка» без сервера:
таблица (Google Таблица, Excel или CSV) - рабочая копия для редакторов,
js/data-records.js - то, что читает сайт. Скрипт переводит одно в другое и
проверяет таблицу перед тем, как она попадёт на сайт.

Выгрузить текущие данные сайта в таблицу (с неё начинается работа):

    python3 tools/data_table.py export усадьбы.xlsx      # или .csv

Загрузить исправленную таблицу обратно. Источник - файл .xlsx/.csv или
ссылка на Google Таблицу (доступ «всем, у кого есть ссылка», на чтение):

    python3 tools/data_table.py import усадьбы.xlsx
    python3 tools/data_table.py import "https://docs.google.com/spreadsheets/d/…/edit#gid=0"

Без --write скрипт только проверяет таблицу и показывает, что изменится
(добавленные, удалённые и изменённые записи). С --write пишет js/data-records.js
(прежняя версия - js/data-records.js.bak) и меняет метку версии этого файла в
index.html, чтобы у посетителей не остался старый кэш.

Ошибки (нет колонки, повтор ID или номера, не число в площади или
координатах, неизвестное поселение) блокируют запись. Предупреждения
(нет контура участка, нет файла скана, необычное значение поля) -
нет, но их стоит просмотреть.
"""
import csv, datetime, io, json, os, re, shutil, sys, urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INDEX = os.path.join(ROOT, "index.html")

# поле записи → заголовок колонки в таблице
COLUMNS = [
    ("id", "ID"), ("num", "№ усадьбы"), ("settlement", "Поселение"), ("street", "Улица / часть"),
    ("buildingType", "Тип постройки"), ("surname", "Фамилия"), ("name", "Имя"), ("patronymic", "Отчество"),
    ("soslovie", "Сословно-профессиональная группа"), ("familyStatus", "Семейное положение"), ("sex", "Пол"),
    ("serviceType", "Род службы"), ("rank", "Чин"), ("position", "Должность"), ("servicePlace", "Место службы"),
    ("registrationPlace", "Место приписки"), ("area_sazh", "Площадь, кв. саж."), ("source", "Источник"),
    ("scanUrl", "Скан (ссылка или путь)"), ("lat", "Широта"), ("lng", "Долгота"),
]
FIELDS = [c[0] for c in COLUMNS]
REQUIRED = ["num", "settlement"]
KNOWN = {
    "sex": {"Мужчины", "Женщины", "Не определено", "Нет данных"},
    "serviceType": {"", "Военная", "Гражданская", "Горная"},
}


# ---------- чтение/запись данных сайта ----------

JS_DIR = os.path.join(ROOT, "js")
# Записи усадеб: js/data-records.js (с 01.10.2026), раньше - js/data.js
RECORDS = os.path.join(JS_DIR, "data-records.js")
if not os.path.exists(RECORDS):
    RECORDS = os.path.join(JS_DIR, "data.js")


def _find_var(text, name):
    m = re.search(r"^var " + name + r" = ", text, re.M)
    if not m:
        return None
    value, end = json.JSONDecoder().raw_decode(text, m.end())
    return value, m.end(), end


def read_var(name, *files):
    """Значение `var <name> = …;` из первого js-файла, где оно есть."""
    for fn in files or ("data-records.js", "data-polygons.js", "data.js"):
        path = os.path.join(JS_DIR, fn)
        if os.path.exists(path):
            found = _find_var(open(path, encoding="utf-8").read(), name)
            if found:
                return found[0]
    return None


def write_records(data):
    """Перезаписывает allData, сохраняя оформление файла (читаемое или в одну строку)."""
    text = open(RECORDS, encoding="utf-8").read()
    _, start, end = _find_var(text, "allData")
    pretty = "\n" in text[start:end]
    body = json.dumps(data, ensure_ascii=False, indent=2) if pretty \
        else json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    shutil.copyfile(RECORDS, RECORDS + ".bak")
    with open(RECORDS, "w", encoding="utf-8", newline="\n") as f:
        f.write(text[:start] + body + text[end:])


def read_datajs():
    parsed = {"allData": read_var("allData")}
    for k in ("ekbPolygons", "niPolygons", "uktPolygons"):
        parsed[k] = read_var(k) or {}
    parsed["COLORS"] = read_var("COLORS") or {}
    return parsed


def polygon_for(rec, polys):
    try:
        key = str(round(float(re.match(r"\d+(\.\d+)?", str(rec["num"])).group(0))))
    except Exception:
        return None
    return polys.get({"Екатеринбург": "ekbPolygons", "Нижне-Исетск": "niPolygons",
                      "Уктус": "uktPolygons"}.get(rec["settlement"]), {}).get(key)


# ---------- таблица ----------

def export(path):
    data = read_datajs()["allData"]
    header = [c[1] for c in COLUMNS]
    rows = [[d.get(f, "") if d.get(f) is not None else "" for f in FIELDS] for d in data]
    if path.lower().endswith(".xlsx"):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Усадьбы"
        ws.append(header)
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F0EBE3")
        widths = {"ID": 7, "№ усадьбы": 10, "Поселение": 14, "Улица / часть": 34, "Источник": 34,
                  "Скан (ссылка или путь)": 38, "Сословно-профессиональная группа": 24}
        for i, h in enumerate(header, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(h, 16)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(path)
    else:
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(header)
            w.writerows(rows)
    print(f"Выгружено записей: {len(rows)} → {path}")


def fetch(src):
    if re.match(r"https?://", src):
        m = re.search(r"/spreadsheets/d/([\w-]+)", src)
        if m:
            gid = re.search(r"[#&?]gid=(\d+)", src)
            src = f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv" + (f"&gid={gid.group(1)}" if gid else "")
        with urllib.request.urlopen(src, timeout=60) as r:
            return "csv", r.read().decode("utf-8-sig")
    return ("xlsx", src) if src.lower().endswith((".xlsx", ".xlsm")) else ("csv", open(src, encoding="utf-8-sig").read())


def read_table(src):
    kind, payload = fetch(src)
    if kind == "xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(payload, read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(it)]
        rows = [["" if v is None else v for v in r] for r in it]
    else:
        first = payload.splitlines()[0] if payload else ""
        delim = max(";,\t", key=first.count)
        rd = list(csv.reader(io.StringIO(payload), delimiter=delim))
        header, rows = [h.strip() for h in rd[0]], rd[1:]
    names = {h.lower(): f for f, h in COLUMNS}
    names.update({f.lower(): f for f in FIELDS})
    colmap = {i: names[h.lower()] for i, h in enumerate(header) if h.lower() in names}
    unknown = [h for i, h in enumerate(header) if i not in colmap and h]
    recs = []
    for n, r in enumerate(rows, start=2):
        if not any(str(v).strip() for v in r):
            continue
        rec = {f: "" for f in FIELDS}
        for i, f in colmap.items():
            if i < len(r):
                rec[f] = r[i]
        rec["_row"] = n
        recs.append(rec)
    return recs, set(colmap.values()), unknown


def num_or_none(v):
    if v is None or str(v).strip() == "":
        return None
    return float(str(v).replace(" ", "").replace(" ", "").replace(",", "."))


def normalize(recs, parsed):
    errors, warns = [], []
    colors = parsed.get("COLORS") or {}
    polys = {k: parsed[k] for k in ("ekbPolygons", "niPolygons", "uktPolygons")}
    old_ids = [d["id"] for d in parsed["allData"]]
    next_id = max(old_ids or [0]) + 1
    seen_id, seen_num, out = {}, {}, []
    for r in recs:
        row = r.pop("_row")
        d = {}
        for f in FIELDS:
            v = r.get(f, "")
            d[f] = v.strip() if isinstance(v, str) else v
        # числа
        for f in ("area_sazh", "lat", "lng"):
            try:
                d[f] = num_or_none(d[f])
            except ValueError:
                errors.append(f"строка {row}: «{d[f]}» в колонке «{dict(COLUMNS)[f]}» - не число")
                d[f] = None
        if isinstance(d["area_sazh"], float) and d["area_sazh"].is_integer():
            d["area_sazh"] = int(d["area_sazh"])
        if d["lat"] is None: d["lat"] = 0.0
        if d["lng"] is None: d["lng"] = 0.0
        # ID
        if d["id"] in ("", None):
            d["id"] = next_id
            next_id += 1
        else:
            try:
                d["id"] = int(float(d["id"]))
            except ValueError:
                errors.append(f"строка {row}: ID «{d['id']}» - не число")
                continue
        if d["id"] in seen_id:
            errors.append(f"строка {row}: ID {d['id']} уже был в строке {seen_id[d['id']]}")
        seen_id[d["id"]] = row
        # номер: 12.0 → 12, как и было в данных
        num = d["num"]
        if isinstance(num, float) and num.is_integer():
            num = str(int(num))
        d["num"] = str(num).strip()
        for f in FIELDS:
            if f not in ("id", "area_sazh", "lat", "lng") and not isinstance(d[f], str):
                d[f] = str(d[f]) if d[f] is not None else ""
        for f in REQUIRED:
            if not d[f]:
                errors.append(f"строка {row}: пустое поле «{dict(COLUMNS)[f]}»")
        if colors and d["settlement"] and d["settlement"] not in colors:
            errors.append(f"строка {row}: неизвестное поселение «{d['settlement']}» (есть: {', '.join(colors)})")
        key = (d["settlement"], d["num"].lower())
        if d["num"] and key in seen_num:
            errors.append(f"строка {row}: {d['settlement']} №{d['num']} уже был в строке {seen_num[key]}")
        seen_num[key] = row
        # предупреждения
        for f, vals in KNOWN.items():
            if d[f] not in vals:
                warns.append(f"строка {row}: необычное значение «{d[f]}» в колонке «{dict(COLUMNS)[f]}»")
        if d["settlement"] and d["num"] and not polygon_for(d, polys) and not (d["lat"] and d["lng"]):
            warns.append(f"строка {row}: {d['settlement']} №{d['num']} - нет контура участка и координат, на карте не появится")
        s = d["scanUrl"]
        if s and not re.match(r"https?://", s):
            if not os.path.isfile(os.path.join(ROOT, s)):
                warns.append(f"строка {row}: файла скана «{s}» нет в папке сайта")
        out.append(d)
    return out, errors, warns


def diff(old, new):
    o = {d["id"]: d for d in old}
    n = {d["id"]: d for d in new}
    added = [n[i] for i in n if i not in o]
    removed = [o[i] for i in o if i not in n]
    changed = []
    for i in n:
        if i in o:
            fields = [f for f in FIELDS if (o[i].get(f) if o[i].get(f) is not None else "") != (n[i].get(f) if n[i].get(f) is not None else "")]
            if fields:
                changed.append((o[i], n[i], fields))
    return added, removed, changed


def label(d):
    who = " ".join(x for x in (d.get("surname"), d.get("name")) if x)
    return f"{d.get('settlement')} №{d.get('num')}" + (f" ({who})" if who else "")


def do_import(src, write):
    parsed = read_datajs()
    old = parsed["allData"]
    recs, cols, unknown = read_table(src)
    missing = [dict(COLUMNS)[f] for f in REQUIRED + ["id"] if f not in cols]
    if missing:
        sys.exit(f"В таблице нет колонок: {', '.join(missing)}")
    if unknown:
        print(f"Колонки, которые сайт не использует (пропущены): {', '.join(unknown)}")
    lost = [dict(COLUMNS)[f] for f in FIELDS if f not in cols]
    if lost:
        print(f"! В таблице нет колонок {', '.join(lost)} - у всех записей эти поля станут пустыми")
    new, errors, warns = normalize(recs, parsed)
    added, removed, changed = diff(old, new)

    print(f"\nЗаписей было {len(old)}, в таблице {len(new)}")
    print(f"Добавлено: {len(added)}, удалено: {len(removed)}, изменено: {len(changed)}")
    for d in added[:15]:
        print(f"  + {label(d)}")
    for d in removed[:15]:
        print(f"  − {label(d)}")
    by_field = {}
    for _, _, fs in changed:
        for f in fs:
            by_field[f] = by_field.get(f, 0) + 1
    if by_field:
        print("  изменённые поля: " + ", ".join(f"{dict(COLUMNS)[f]} - {n}" for f, n in sorted(by_field.items(), key=lambda x: -x[1])))
    for o, n, fs in changed[:15]:
        print(f"  ~ {label(n)}: " + "; ".join(f"{dict(COLUMNS)[f]}: «{o.get(f) or ''}» → «{n.get(f) or ''}»" for f in fs[:4]))
    if len(removed) > 20:
        print(f"! Удаляется много записей ({len(removed)}). Проверьте, что выгружен нужный лист таблицы.")

    if warns:
        print(f"\nПредупреждения ({len(warns)}):")
        for w in warns[:40]:
            print("  · " + w)
        if len(warns) > 40:
            print(f"  … и ещё {len(warns) - 40}")
    if errors:
        print(f"\nОшибки ({len(errors)}) - данные не записаны:")
        for e in errors[:60]:
            print("  ✗ " + e)
        sys.exit(1)

    if not (added or removed or changed):
        print("\nИзменений нет.")
        return
    if not write:
        print("\nПроверка пройдена. Чтобы записать, добавьте --write")
        return

    # порядок полей - как в текущем файле данных, чтобы в диффе были видны только правки
    old_order = {d["id"]: list(d.keys()) for d in old}
    default = list(old[0].keys()) if old else FIELDS
    def ordered(d):
        order = old_order.get(d["id"], default)
        return {**{k: d[k] for k in order if k in d}, **{k: v for k, v in d.items() if k not in order}}
    new = [ordered(d) for d in new]
    write_records(new)
    stamp = datetime.datetime.now().strftime("%Y%m%d%H%M")
    rel = os.path.relpath(RECORDS, ROOT).replace(os.sep, "/")
    html = open(INDEX, encoding="utf-8").read()
    html2 = re.sub(r'src="' + re.escape(rel) + r'(\?v=[^"]*)?"', f'src="{rel}?v={stamp}"', html)
    if html2 != html:
        open(INDEX, "w", encoding="utf-8", newline="\n").write(html2)
    print(f"\n{rel} записан (прежняя версия - {rel}.bak), метка версии в index.html: {stamp}.")
    print(f"Выложите на хостинг {rel} и index.html.")


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 2 or args[0] not in ("export", "import"):
        print(__doc__)
        sys.exit(2)
    if args[0] == "export":
        export(args[1])
    else:
        do_import(args[1], "--write" in sys.argv)


if __name__ == "__main__":
    main()
